"""Build a resume DOCX snapshot and an Anthology Fellows CV from shared facts.

Run: python scripts/build_documents.py
Dependencies: python-docx, reportlab, pypdf. Local .tools/python is supported.
The editable resume source is resumes/Hrudayangam-Mehta-Resume.tex.
Build its PDF/text with scripts/build_latex_resume.py. This script never overwrites
the LaTeX source or its compiled PDF/text. The resume DOCX is a separate snapshot
of the shared profile, so later edits to the LaTeX file do not update that DOCX.
The Anthology CV is generated as PDF, DOCX, and text in resumes/.
"""

from __future__ import annotations

import html
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".tools" / "python"))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, HRFlowable


@dataclass
class Block:
    kind: str
    text: str = ""


def esc(value):
    return html.escape(str(value), quote=True)


def link(text, url):
    return f'<a href="{esc(url)}">{esc(text)}</a>'


def bold(value):
    return f"<b>{esc(value)}</b>"


def strip_markup(text):
    return html.unescape(re.sub(r"<[^>]+>", "", text))


def header(profile, cv=False):
    contact = " | ".join([
        esc(profile["location"]),
        link(profile["email"], "mailto:" + profile["email"]),
        esc(profile["phone"]),
    ])
    labels = {item["label"]: item["url"] for item in profile["links"]}
    online = " | ".join([
        link("github.com/Hrudayangam", labels["GitHub"]),
        link("linkedin.com/in/hrudaymehta", labels["LinkedIn"]),
        link("hrudaymehta.vercel.app", labels["Website"]),
    ])
    blocks = [Block("name", esc(profile["name"])),
              Block("contact", contact), Block("contact", online)]
    if cv:
        blocks.append(Block("tagline", "Machine Learning Research | Anthology Fellows application"))
        blocks.append(Block("body", esc(profile["anthology_summary"])))
    return blocks


def experience_blocks(profile, include_early=False, concise=False):
    blocks = [Block("section", "RESEARCH EXPERIENCE")]
    for item in profile["experience"]:
        if not item.get("include_in_documents", True):
            continue
        if not include_early and "Intern" in item["title"]:
            continue
        organization = item["organization"]
        if concise:
            organization = organization.replace(
                "Autonomous Intelligent Robotics (AIR) Group, Binghamton University",
                "AIR Group, Binghamton University",
            )
        blocks.append(Block("entry", f'{bold(item["title"])} | {esc(organization)}'))
        if item.get("dates"):
            blocks.append(Block("meta", esc(item["dates"])))
        bullets = item["bullets"]
        for bullet in bullets:
            blocks.append(Block("bullet", esc(bullet)))
    return blocks


def skill_blocks(profile):
    blocks = [Block("section", "TECHNICAL SKILLS")]
    for group in profile["skills"]:
        blocks.append(Block("body", f'{bold(group["label"] + ":")} {esc(", ".join(group["items"]))}'))
    return blocks


def education_blocks(profile, concise=False):
    blocks = [Block("section", "EDUCATION")]
    for item in profile["education"]:
        degree = item["degree"]
        institution = item["institution"]
        if concise:
            degree = degree.replace("Computer Science and Engineering", "Computer Science & Engineering")
            institution = institution.replace(", SUNY", "").replace("Vellore Institute of Technology, Bhopal", "VIT Bhopal")
        blocks.append(Block("entry", f'{bold(degree)} | {esc(institution)}'))
        detail = item.get("detail", "")
        if concise:
            if "PhD" in degree:
                detail = "Adviser: " + profile["advisor"]["name"]
            elif "MS" in degree:
                detail = "GPA: 3.83/4.00"
            else:
                detail = ""
        metadata = " | ".join(x for x in [item["dates"], detail] if x)
        blocks.append(Block("meta", esc(metadata)))
    return blocks


def patent_blocks(evidence, concise=False):
    blocks = [Block("section", "PATENT APPLICATION")]
    for patent in evidence["patents"]:
        blocks.append(Block("entry", link(patent["title"], patent["url"])))
        if not concise:
            blocks.append(Block("body", esc(", ".join(patent["inventors"])) + "."))
        blocks.append(Block("body", esc("Co-inventor. " + patent["recommended_label"] + ".")))
    return blocks


def make_resume(profile, evidence):
    blocks = header(profile)
    blocks.append(Block("tagline", "PhD Researcher | LLM Evaluation | AI Alignment"))
    blocks += experience_blocks(profile, concise=True)
    blocks.append(Block("section", "SELECTED RESEARCH AND PUBLICATIONS"))
    publications = {item["id"]: item for item in evidence["publications"]}
    selection = ["antisemitism-emnlp-2025", "floorplan-journal-2025", "woofs-words-2026"]
    for key in selection:
        paper = publications[key]
        blocks.append(Block("entry", link(paper["title"], paper["url"])))
        blocks.append(Block("meta", esc(paper["venue_short"])))
        blocks.append(Block("body", esc(paper["contribution"])))
    blocks += skill_blocks(profile)
    blocks += education_blocks(profile, concise=True)
    return blocks


def make_cv(profile, evidence):
    blocks = header(profile, cv=True)
    blocks += experience_blocks(profile)
    blocks.append(Block("section", "SELECTED TECHNICAL RESEARCH"))
    for project in profile["projects"]:
        blocks.append(Block("entry", link(project["name"], project["url"])))
        blocks.append(Block("meta", esc(project["description"])))
        for bullet in project["bullets"]:
            blocks.append(Block("bullet", esc(bullet)))
    blocks.append(Block("pagebreak"))
    blocks.append(Block("section", "SELECTED PUBLICATIONS"))
    blocks.append(Block("meta", "* Equal contribution."))
    for paper in evidence["publications"]:
        blocks.append(Block("entry", link(paper["title"], paper["url"])))
        equal = set(paper.get("equal_contribution", []))
        authors = []
        for author in paper["authors"]:
            label = author + ("*" if author in equal else "")
            authors.append(bold(label) if author == profile["name"] else esc(label))
        blocks.append(Block("body", ", ".join(authors) + "."))
        journal = paper["venue"]
        if paper.get("volume"):
            journal += f', {paper["volume"]}({paper["issue"]})'
        if paper.get("pages"):
            journal += f': {paper["pages"]}'
        elif paper.get("article_number"):
            journal += f', article {paper["article_number"]}'
        journal += f'. {paper["year"]}.'
        blocks.append(Block("body", esc(journal)))
        blocks.append(Block("body", bold("Contribution: ") + esc(paper["contribution"])))
        links = [link("DOI: " + paper["doi"], "https://doi.org/" + paper["doi"])]
        if paper.get("code"):
            links.append(link("Code", paper["code"]))
        if paper.get("project"):
            links.append(link("Project", paper["project"]))
        blocks.append(Block("body", " | ".join(links)))
        if paper["id"] == "floorplan-journal-2025":
            alternate = next((u for u in evidence.get("unresolved", []) if "openreview.net/forum" in u["url"]), None)
            if alternate:
                blocks.append(Block("meta", link("Additional accepted version: OpenReview", alternate["url"]) + ". Venue/year: TBD."))
    blocks += patent_blocks(evidence)
    blocks += education_blocks(profile)
    blocks += skill_blocks(profile)
    internships = [item for item in profile["experience"] if "Intern" in item["title"] and item.get("include_in_documents", True)]
    if internships:
        blocks.append(Block("section", "EARLIER EXPERIENCE"))
        for item in internships:
            blocks.append(Block("entry", f'{bold(item["title"])} | {esc(item["organization"])} | {esc(item["dates"])}'))
            for bullet in item["bullets"]:
                blocks.append(Block("body", esc(bullet)))
    return blocks


def register_fonts():
    candidates = [Path("C:/Windows/Fonts"), Path("/usr/share/fonts/truetype/msttcorefonts")]
    for directory in candidates:
        variants = {"Resume": "arial.ttf", "Resume-Bold": "arialbd.ttf", "Resume-Italic": "ariali.ttf", "Resume-BoldItalic": "arialbi.ttf"}
        if all((directory / name).exists() for name in variants.values()):
            for name, filename in variants.items():
                pdfmetrics.registerFont(TTFont(name, str(directory / filename)))
            pdfmetrics.registerFontFamily("Resume", normal="Resume", bold="Resume-Bold", italic="Resume-Italic", boldItalic="Resume-BoldItalic")
            return "Resume"
    return "Helvetica"


def pdf_styles(cv=False):
    font = register_fonts()
    bold_font = font + "-Bold"
    size = 10.1 if cv else 10.2
    common = dict(fontName=font, fontSize=size, leading=size * 1.23, textColor=colors.HexColor("#171717"), alignment=TA_LEFT, allowWidows=0, allowOrphans=0)
    def style(name, **kwargs):
        return ParagraphStyle(name, **(common | kwargs))
    return {
        "name": style("name", fontName=bold_font, fontSize=21, leading=25, spaceAfter=4),
        "contact": style("contact", fontSize=9, leading=11.5, spaceAfter=1),
        "tagline": style("tagline", fontName=bold_font, fontSize=10, leading=13, spaceBefore=6, spaceAfter=4),
        "section": style("section", fontName=bold_font, fontSize=10, leading=12, spaceBefore=10 if cv else 8, spaceAfter=5, keepWithNext=True),
        "entry": style("entry", fontName=bold_font, spaceBefore=4 if cv else 3, spaceAfter=2, keepWithNext=True),
        "meta": style("meta", fontSize=size - 0.45, leading=size * 1.2, textColor=colors.HexColor("#444444"), spaceAfter=3, keepWithNext=True),
        "body": style("body", spaceAfter=4 if cv else 2),
        "bullet": style("bullet", leftIndent=10, firstLineIndent=-8, spaceAfter=3 if cv else 2),
    }


def write_pdf(path, blocks, profile, cv=False):
    styles = pdf_styles(cv)
    story = []
    for block in blocks:
        if block.kind == "pagebreak":
            story.append(PageBreak())
            continue
        text = block.text
        if block.kind == "bullet":
            text = "&#8226; " + text
        story.append(Paragraph(text, styles[block.kind]))
        if block.kind == "section":
            story.append(HRFlowable(width="100%", thickness=0.45, color=colors.HexColor("#777777"), spaceAfter=2))
    title = profile["name"] + (" - Anthology Fellows CV" if cv else " - Resume")
    document = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=39.6, leftMargin=39.6,
                                 topMargin=31.5, bottomMargin=33, title=title, author=profile["name"])
    def footer(canvas, doc):
        if cv:
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#555555"))
            canvas.drawString(39.6, 19, profile["name"])
            canvas.drawRightString(letter[0] - 39.6, 19, str(doc.page))
            canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def add_word_markup(paragraph, markup):
    root = ElementTree.fromstring("<root>" + markup + "</root>")
    def append_text(text, strong=False, italic=False, url=None):
        if not text:
            return
        if url:
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            if strong:
                props.append(OxmlElement("w:b"))
            if italic:
                props.append(OxmlElement("w:i"))
            run.append(props)
            content = OxmlElement("w:t")
            content.set(qn("xml:space"), "preserve")
            content.text = text
            run.append(content)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            run = paragraph.add_run(text)
            if strong:
                run.bold = True
            if italic:
                run.italic = True
    def walk(element, strong=False, italic=False, url=None):
        strong = strong or element.tag == "b"
        italic = italic or element.tag == "i"
        url = element.attrib.get("href", url)
        append_text(element.text, strong, italic, url)
        for child in element:
            walk(child, strong, italic, url)
            append_text(child.tail, strong, italic, url)
    walk(root)


def write_docx(path, blocks, profile, cv=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.44)
    section.bottom_margin = Inches(0.46)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.18)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.1 if cv else 10.2)
    normal.font.color.rgb = RGBColor.from_string("171717")
    normal.paragraph_format.space_after = Pt(4 if cv else 2)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True
    for block in blocks:
        if block.kind == "pagebreak":
            doc.add_page_break()
            continue
        paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(0)
        if block.kind == "name":
            fmt.space_after = Pt(4)
        elif block.kind == "contact":
            fmt.space_after = Pt(1)
        elif block.kind == "tagline":
            fmt.space_before = Pt(6)
            fmt.space_after = Pt(4)
        elif block.kind == "section":
            fmt.space_before = Pt(10 if cv else 8)
            fmt.space_after = Pt(5)
            fmt.keep_with_next = True
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in {"val": "single", "sz": "4", "space": "3", "color": "777777"}.items():
                bottom.set(qn("w:" + key), value)
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
        elif block.kind == "entry":
            fmt.space_before = Pt(4 if cv else 3)
            fmt.space_after = Pt(2)
            fmt.keep_with_next = True
        elif block.kind == "meta":
            fmt.space_after = Pt(3)
            fmt.keep_with_next = True
        elif block.kind == "bullet":
            fmt.left_indent = Pt(10)
            fmt.first_line_indent = Pt(-8)
            fmt.space_after = Pt(3 if cv else 2)
            paragraph.add_run("\u2022 ")
        add_word_markup(paragraph, block.text)
        sizes = {"name": 21, "contact": 9, "tagline": 10, "section": 10, "meta": (9.65 if cv else 9.75)}
        # Apply formatting to all runs, including hyperlinks.
        for run in paragraph._p.findall(".//" + qn("w:r")):
            props = run.find(qn("w:rPr"))
            if props is None:
                props = OxmlElement("w:rPr")
                run.insert(0, props)
            if block.kind in sizes:
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), str(round(sizes[block.kind] * 2)))
                props.append(size)
            if block.kind in {"name", "tagline", "section", "entry"}:
                props.append(OxmlElement("w:b"))
    doc.core_properties.author = profile["name"]
    doc.core_properties.title = profile["name"] + (" - Anthology Fellows CV" if cv else " - Resume")
    doc.core_properties.subject = "Machine learning research, evaluation, and data engineering"
    if cv:
        footer = section.footer.paragraphs[0]
        footer.add_run(profile["name"] + " | ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        for run in footer.runs:
            run.font.size = Pt(8)
    doc.save(path)


def verify(pdf_path, docx_path, expected_pages, blocks):
    reader = PdfReader(pdf_path)
    if len(reader.pages) != expected_pages:
        raise RuntimeError(f"{pdf_path.name}: expected {expected_pages} pages, found {len(reader.pages)}")
    extracted = "\n".join(page.extract_text() for page in reader.pages)
    for term in ["Hrudayangam Mehta", "EMNLP", "63/884,317", "Jeremy Blackburn"]:
        if term not in extracted:
            raise RuntimeError(f"Missing searchable PDF text: {term}")
    for term in ["preprint", "Intrusion Detection", "Instruction Pipeline Simulator", "enthusiast", "iSmriti"]:
        if term.lower() in extracted.lower():
            raise RuntimeError(f"Unexpected legacy text: {term}")
    doc = Document(docx_path)
    if doc.tables:
        raise RuntimeError("Resume/CV must not use layout tables")
    docx_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    normalize = lambda text: re.sub(r"\s+", "", text)
    for block in blocks:
        expected = normalize(strip_markup(block.text))
        if expected and expected not in normalize(extracted):
            raise RuntimeError(f"Missing PDF content: {strip_markup(block.text)}")
        if expected and expected not in normalize(docx_text):
            raise RuntimeError(f"Missing DOCX content: {strip_markup(block.text)}")
    links = sum(len(page.get("/Annots", [])) for page in reader.pages)
    return {"pdf": pdf_path.name, "docx": docx_path.name, "pages": len(reader.pages), "pdf_annotations": links,
            "text_characters": len(extracted), "docx_tables": len(doc.tables), "all_content_verified": True}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    profile = json.loads((ROOT / "content/profile.json").read_text(encoding="utf-8-sig"))
    evidence = json.loads((ROOT / "content/research-evidence.json").read_text(encoding="utf-8-sig"))
    output = ROOT / "resumes"
    output.mkdir(exist_ok=True)
    report = []
    # Keep the hand-edited .tex and its PDF/text authoritative for the resume.
    resume_blocks = make_resume(profile, evidence)
    resume_docx = output / "Hrudayangam-Mehta-Resume.docx"
    write_docx(resume_docx, resume_blocks, profile)
    doc = Document(resume_docx)
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    if doc.tables or "ismriti" in doc_text.lower():
        raise RuntimeError("Unexpected layout table or excluded experience in resume DOCX")
    for block in resume_blocks:
        if re.sub(r"\s+", "", strip_markup(block.text)) not in re.sub(r"\s+", "", doc_text):
            raise RuntimeError(f"Missing DOCX content: {strip_markup(block.text)}")
    report.append({"docx": resume_docx.name, "source": "Shared-profile snapshot; LaTeX resume PDF is built separately", "docx_tables": 0, "all_content_verified": True})
    for slug, blocks, cv, pages in [
        ("Hrudayangam-Mehta-Anthology-CV", make_cv(profile, evidence), True, 2),
    ]:
        write_pdf(output / (slug + ".pdf"), blocks, profile, cv)
        write_docx(output / (slug + ".docx"), blocks, profile, cv)
        text = "\n".join(("- " if b.kind == "bullet" else "") + strip_markup(b.text) if b.kind != "pagebreak" else "\n" for b in blocks)
        (output / (slug + ".txt")).write_text(text + "\n", encoding="utf-8")
        report.append(verify(output / (slug + ".pdf"), output / (slug + ".docx"), pages, blocks))
    (output / "validation.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
